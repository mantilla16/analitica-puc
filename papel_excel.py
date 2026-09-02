"""
El papel de trabajo, en Excel.

Es el formato en el que un papel se archiva y se revisa de verdad, así
que la regla que ordena este módulo es una sola: **los números van como
números**, con formato de celda, no como texto ya formateado. Un papel
donde el auditor no puede sumar una columna ni filtrar una tabla es una
imagen del papel, no el papel.

Consume exactamente el mismo dict que arma papel.py. No recalcula nada:
si esta hoja mostrara una cifra distinta a la de la pantalla, el papel
dejaría de ser reproducible.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Paleta de marca, en el formato que entiende Excel
NAVY = "1B1A5E"
CIAN = "E2F7FC"
GRIS = "E3E7F2"
VERDE = "DDF3EA"
ROJO = "FDEAE7"
AMBAR = "FDF0DA"

PESOS = '#,##0.00'
PORC = '0.00"%"'

_titulo = Font(name="Calibri", size=14, bold=True, color=NAVY)
_encabezado = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
_rotulo = Font(name="Calibri", size=9, bold=True, color="5A5A7A")
_normal = Font(name="Calibri", size=10)
_relleno_enc = PatternFill("solid", fgColor=NAVY)
_borde = Border(*[Side(style="thin", color="D8DDE8")] * 4)

TONO_ESTADO = {"OK": VERDE, "ALERTA": AMBAR, "NO_EJECUTADO": AMBAR,
               "FALLA": ROJO, "BLOQUEANTE": ROJO}


def _num(v):
    """Decimal/str -> float para que Excel lo trate como número. Devuelve
    el original si no es convertible: es preferible una celda de texto a
    perder el dato."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return float(v)
    try:
        return float(str(v).replace(".", "").replace(",", "."))
    except ValueError:
        return v


def _hoja(wb, nombre, titulo, subtitulo=None):
    h = wb.create_sheet(nombre[:31])
    h["A1"] = titulo
    h["A1"].font = _titulo
    if subtitulo:
        h["A2"] = subtitulo
        h["A2"].font = _rotulo
        h["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    h.sheet_view.showGridLines = False
    return h


def _fecha(v) -> str:
    """Fecha legible. Las marcas de tiempo llegan como datetime desde la base,
    pero un papel restaurado de un JSON las trae como texto."""
    if not v:
        return ""
    return (v.strftime("%d/%m/%Y %H:%M") if hasattr(v, "strftime")
            else str(v)[:16].replace("T", " "))


def _tabla(h, fila_ini, columnas, filas, anchos=None, formatos=None,
           principal=False):
    """Escribe una tabla con encabezado.

    `principal` marca la tabla que manda en la hoja: solo esa recibe el
    panel congelado y el autofiltro, porque Excel admite UNO por hoja. Si
    se fijan en cada tabla, gana la última -- y en una hoja con varias
    tablas eso deja el panel congelado en una fila muy abajo, con lo que
    Excel congela toda el área visible y el scroll deja de responder."""
    for j, c in enumerate(columnas, start=1):
        celda = h.cell(row=fila_ini, column=j, value=c)
        celda.font = _encabezado
        celda.fill = _relleno_enc
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = _borde

    for i, f in enumerate(filas, start=fila_ini + 1):
        for j, v in enumerate(f, start=1):
            celda = h.cell(row=i, column=j, value=v)
            celda.font = _normal
            celda.border = _borde
            if formatos and formatos.get(j):
                celda.number_format = formatos[j]
                celda.alignment = Alignment(horizontal="right")

    for j, a in enumerate(anchos or [], start=1):
        h.column_dimensions[get_column_letter(j)].width = a

    fin = fila_ini + len(filas)
    if filas and principal:
        h.auto_filter.ref = f"A{fila_ini}:{get_column_letter(len(columnas))}{fin}"
        h.freeze_panes = h.cell(row=fila_ini + 1, column=1)
    return fin + 2


def _bloque_texto(h, fila, etiqueta, texto, ancho=110):
    h.cell(row=fila, column=1, value=etiqueta).font = _rotulo
    c = h.cell(row=fila + 1, column=1, value=texto)
    c.font = _normal
    c.alignment = Alignment(wrap_text=True, vertical="top")
    h.merge_cells(start_row=fila + 1, start_column=1, end_row=fila + 3, end_column=8)
    h.column_dimensions["A"].width = min(ancho, 60)
    return fila + 5


# =====================================================================

def construir(p: dict) -> bytes:
    """Recibe el papel ya armado y devuelve el .xlsx en memoria."""
    wb = Workbook()
    wb.remove(wb.active)

    id_ = p["identificacion"]
    d = p["comparativo"]

    # ------------------------------------------------------- 1. portada
    h = _hoja(wb, "Papel", id_["papel"], id_["norma"])
    fila = 4
    for et, v in [
        ("Cliente", id_["cliente"]), ("NIT", id_["nit"]),
        ("Fecha de corte", id_["fecha_corte"]),
        ("Comparativo clases 1-3", id_["cierre_anterior"]),
        ("Comparativo clases 4-7", id_["corte_anterior"]),
        ("Fase", id_["fase"]), ("Responsable", id_["responsable"] or "—"),
        ("Materialidad aplicada", _num(d.get("umbral"))),
        ("Generado", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Versión del papel", p["version"]),
    ]:
        h.cell(row=fila, column=1, value=et).font = _rotulo
        c = h.cell(row=fila, column=2, value=v)
        c.font = _normal
        if et == "Materialidad aplicada" and isinstance(v, float):
            c.number_format = PESOS
        fila += 1

    h.column_dimensions["A"].width = 26
    h.column_dimensions["B"].width = 46

    fila += 1
    con = p["conclusion"]
    h.cell(row=fila, column=1, value="CONCLUSIÓN").font = _titulo
    h.cell(row=fila, column=2, value=con["estado"].replace("_", " ")).font = Font(
        name="Calibri", size=12, bold=True,
        color={"RAZONABLE": "0D7A57", "NO_CONCLUYENTE": "C0392B"}.get(
            con["estado"], "B56A00"))
    fila += 1
    h.cell(row=fila, column=1, value="Riesgo").font = _rotulo
    h.cell(row=fila, column=2,
           value=f"{p['riesgo']['nivel']} ({p['riesgo']['puntos']} puntos)").font = _normal
    # Un párrafo por fila, no todo en una celda combinada: así se puede leer
    # sin ampliar la fila a media pantalla, y se puede copiar suelto al
    # informe. El texto es el mismo que muestra la pantalla.
    fila += 2
    h.cell(row=fila, column=1, value="Procedimiento y fundamento").font = _rotulo
    fila += 1
    for par in con.get("parrafos") or [con["texto"]]:
        c = h.cell(row=fila, column=1, value=par)
        c.font = _normal
        c.alignment = Alignment(wrap_text=True, vertical="top")
        h.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=8)
        h.row_dimensions[fila].height = max(30, 13 * (len(par) // 110 + 1))
        fila += 1
    fila += 1

    # ------------------------------------------- 2. contrato de datos
    h = _hoja(wb, "Contrato de datos",
              "Contrato de datos",
              "De qué archivo y de qué columna salió cada cifra. La huella SHA-256 "
              "ata este papel a un archivo concreto.")
    fila = _tabla(h, 4,
        ["Insumo", "Archivo", "Huella SHA-256", "Hoja", "Periodo",
         "Filas leídas", "Filas promovidas", "Estado", "Cargado por"],
        [[c["insumo"],
          c.get("archivo", "— sin cargar —"),
          c.get("huella_sha256", ""), c.get("hoja", ""),
          f"{c['periodo']['inicio']} a {c['periodo']['fin']}" if c.get("cargado") else "",
          c.get("filas_leidas"), c.get("filas_promovidas"),
          c.get("estado", "SIN CARGAR"), c.get("subido_por", "")]
         for c in p["contrato_datos"]],
        anchos=[34, 46, 34, 12, 26, 12, 14, 14, 14], principal=True)

    # el mapeo campo a campo va debajo, por insumo
    for c in p["contrato_datos"]:
        if not c.get("columnas"):
            continue
        h.cell(row=fila, column=1,
               value=f"Origen de cada campo — {c['insumo']}").font = _rotulo
        fila = _tabla(h, fila + 1, ["Campo del estándar", "Columna del archivo"],
                      [[k, v] for k, v in c["columnas"].items()],
                      anchos=[34, 46])

    # ----------------------------------------------- 3. controles
    h = _hoja(wb, "Controles", "Controles ejecutados",
              "Un control solo es evidencia si puede fallar y si contrasta contra "
              "algo que no se derive de lo que verifica.")
    filas = []
    for c in p["controles_previos"] + p["gates"]:
        filas.append([c["codigo"], c["marca"], c["nombre"], c["estado"],
                      "Sí" if c["es_evidencia"] else "No — consistencia interna",
                      c["detalle"]])
    fin = _tabla(h, 4,
        ["Código", "Marca", "Control", "Estado", "¿Evidencia?", "Detalle"],
        filas, anchos=[9, 7, 34, 15, 26, 90], principal=True)

    # color por estado, para que el resultado se lea sin recorrer el texto
    for i, c in enumerate(p["controles_previos"] + p["gates"], start=5):
        tono = TONO_ESTADO.get(c["estado"])
        if tono:
            h.cell(row=i, column=4).fill = PatternFill("solid", fgColor=tono)
        h.cell(row=i, column=6).alignment = Alignment(wrap_text=True, vertical="top")

    # ------------------------------------------------ 3.b descuadres
    # Los controles que fallaron, cuenta por cuenta. Una hoja que diga
    # "falló" sin decir dónde obliga a salir del papel a buscarlo.
    INSUMO = {"BAL_ACTUAL": "Balance a la fecha de corte",
              "BAL_CIERRE_ANTERIOR": "Balance al 31-dic del año anterior",
              "BAL_CORTE_ANTERIOR": "Balance al mismo corte del año anterior"}
    # Se recorre por nombre de clave y no adivinando la forma de la lista:
    # `cifras` fue creciendo (redondeos, cuentas afectadas, documentos) y
    # cualquier lista de diccionarios pasaba por descuadre de línea.
    detalle = []
    for c in p["gates"] + p["controles_previos"]:
        cifras = c.get("cifras") or {}
        falla = c["estado"] in ("FALLA", "BLOQUEANTE")

        if falla:
            for tipo, filas in cifras.items():
                if not (isinstance(filas, list) and filas
                        and isinstance(filas[0], dict)
                        and "codigo" in filas[0]):
                    continue
                for f in filas:                       # descuadre de línea
                    detalle.append([
                        c["codigo"], "Falla", INSUMO.get(tipo, tipo),
                        f["codigo"], f["nombre"],
                        "saldo inicial + débitos − créditos ≠ saldo final",
                        _num(f["diferencia"])])

            for f in cifras.get("no_cuadran") or []:  # movimientos vs balance
                detalle.append([
                    c["codigo"], "Falla", "Movimientos del periodo",
                    f["cuenta"], f["nombre"],
                    f"neto {f['neto_movimientos']} contra {f['contra']} "
                    f"{f['esperado']}",
                    _num(f["diferencia"])])

            for x in cifras.get("documentos_descuadrados") or []:
                detalle.append([
                    c["codigo"], "Falla", "Asiento del archivo de movimientos",
                    x["num_doc"],
                    f"{x['lineas']} líneas · {_fecha(x['fecha'])}",
                    "los débitos del documento no igualan sus créditos: "
                    f"débito {x['debito']} contra crédito {x['credito']}",
                    _num(x["diferencia"])])

            for a in cifras.get("cuentas_afectadas") or []:
                detalle.append([
                    c["codigo"], "Cuenta alcanzada",
                    "Asientos que no cuadran", a["cuenta"],
                    f"{a['lineas']} líneas",
                    "el cruce contra movimientos no concluye en esta cuenta: "
                    "parte de sus líneas vienen de asientos incompletos",
                    _num(a["monto"])])

        # Las cuentas que pasaron USANDO la tolerancia van en la misma hoja:
        # una tolerancia que se aplica y no se declara equivale a no haber
        # corrido el control.
        for f in cifras.get("por_redondeo") or []:
            detalle.append([
                c["codigo"], "Dentro de tolerancia",
                "Movimientos del periodo", f["cuenta"], f["nombre"],
                f"neto {f['neto_movimientos']} contra {f['contra']} "
                f"{f['esperado']} — diferencia admitida como redondeo "
                f"(tolerancia {cifras.get('tolerancia_por_cuenta')} por cuenta)",
                _num(f["diferencia"])])

    if detalle:
        h = _hoja(wb, "Descuadres", "Descuadres y diferencias toleradas",
                  "Las cuentas y documentos concretos detrás de cada control que "
                  "falló, y lo que pasó usando la tolerancia declarada. Es por "
                  "donde empieza la revisión.")
        _tabla(h, 4, ["Control", "Estado", "Fuente", "Cuenta / documento",
                      "Nombre", "Detalle", "Diferencia"],
               detalle, anchos=[10, 20, 34, 18, 40, 68, 20],
               formatos={7: PESOS}, principal=True)

    # ------------------------------------------- 4. cédula sumaria
    h = _hoja(wb, "Cédula sumaria", "Cédula sumaria — saldos por clase",
              "Saldo en naturaleza al corte, a nivel de cuenta.")
    _tabla(h, 4, ["Clase", "Nombre", "Cuentas", "Saldo"],
           [[c["clase"], c["clase_nombre"], c["cuentas"], _num(c["saldo"])]
            for c in p["cedula_sumaria"]],
           anchos=[9, 38, 12, 22], formatos={4: PESOS}, principal=True)

    # ------------------------------------------------ 5. comparativo
    h = _hoja(wb, "Comparativo", "Comparativo por cuenta (NIA 520)",
              "Clases 1 a 3 contra el cierre anterior; 4 a 7 contra el mismo corte "
              "del año anterior. Marca Δ.")
    _tabla(h, 4,
        ["Cuenta", "Nombre", "Clase", "Comparado contra", "Saldo actual",
         "Saldo comparativo", "Variación", "%", "Motivo", "Marca"],
        [[f["cuenta"], f["nombre"], f["clase"], f["regla"],
          _num(f["saldo_actual"]), _num(f["saldo_comparativo"]),
          _num(f["variacion"]), _num(f["variacion_pct"]),
          f["motivo"] or "", "Δ" if f["significativa"] else ""]
         for f in d["filas"]],
        anchos=[11, 40, 8, 30, 20, 20, 20, 10, 18, 8],
        formatos={5: PESOS, 6: PESOS, 7: PESOS, 8: PORC}, principal=True)

    # ------------------------------------------- 5.b análisis del modelo
    # Hoja aparte y no una columna del comparativo: el texto es largo y en
    # una celda al lado de las cifras no se lee. Va con su procedencia --
    # modelo, versión, si el auditor lo instruyó, si sus cifras se
    # contrastaron contra el motor -- porque un texto de IA sin trazabilidad
    # no es papel de trabajo.
    obs = p.get("observaciones") or {}
    analisis = []
    for f in d["filas"]:
        o = obs.get(f["cuenta"])
        if not f["significativa"] or not o:
            continue
        sin_ver = o.get("cifras_no_verificadas") or []
        analisis.append([
            f["cuenta"], f["nombre"], _num(f["variacion"]), f["motivo"] or "",
            o["texto"],
            "Sí" if o.get("verificado") else "No",
            ", ".join(str(x) for x in sin_ver),
            o.get("modelo") or "", o.get("version"),
            o.get("instruccion_auditor") or "",
            o.get("creado_por") or "", _fecha(o.get("creado_en")),
        ])
    if analisis:
        h = _hoja(wb, "Análisis", "Análisis de las variaciones",
                  "Redactado por el modelo a partir de las cifras del motor y de "
                  "los movimientos de cada cuenta. El modelo explica y escribe; "
                  "no calcula. Las cifras que no se pudieron contrastar contra el "
                  "motor se listan para revisión.")
        _tabla(h, 4,
               ["Cuenta", "Nombre", "Variación", "Motivo", "Análisis",
                "¿Cifras verificadas?", "Cifras sin verificar", "Modelo",
                "Versión", "Instrucción del auditor", "Generado por", "Fecha"],
               analisis,
               anchos=[11, 34, 20, 16, 110, 12, 26, 22, 9, 44, 18, 12],
               formatos={3: PESOS}, principal=True)
        for fila in range(5, 5 + len(analisis)):
            h.cell(row=fila, column=5).alignment = Alignment(
                wrap_text=True, vertical="top")
            h.cell(row=fila, column=10).alignment = Alignment(
                wrap_text=True, vertical="top")

    # ------------------------------------- 5.c cuentas fuera de catálogo
    fuera = next((c["cifras"].get("fuera_de_catalogo")
                  for c in p["controles_previos"] if c["codigo"] == "P04"), None)
    if fuera:
        h = _hoja(wb, "Fuera de catálogo", "Cuentas fuera del catálogo PUC",
                  "No existen en el Decreto 2650. Puede ser normal, pero su "
                  "naturaleza se heredó del prefijo y de ella depende el signo con "
                  "que la cuenta entra al comparativo.")
        _tabla(h, 4,
               ["Cuenta", "Nombre en el balance", "Clase",
                "Naturaleza supuesta", "¿De dónde salió?",
                "¿El saldo la contradice?", "Vecinos en el catálogo",
                "Saldo final", "Saldo natural"],
               [[x["codigo_puc"], x["nombre_cuenta"], x["clase"],
                 "débito" if x["signo"] == 1 else "crédito",
                 (f"declarada en {x['prefijo_naturaleza']}"
                  if x.get("naturaleza_declarada")
                  else f"heredada de la clase {x['clase']}"),
                 "SÍ — revisar" if x.get("saldo_contradice_naturaleza") else "no",
                 " · ".join(
                     f"{v['codigo']} {v['nombre']}"
                     + (f" ({'débito' if v['naturaleza'] == 'D' else 'crédito'})"
                        if v.get("naturaleza") else "")
                     for v in x.get("vecinos_catalogo") or []),
                 _num(x["saldo_final"]), _num(x["saldo_natural"])]
                for x in fuera],
               anchos=[11, 44, 8, 20, 30, 24, 64, 20, 20],
               formatos={8: PESOS, 9: PESOS}, principal=True)

    # ---------------------------------------------------- 6. alcance
    h = _hoja(wb, "Alcance", "Alcance y selección")
    fila = 4
    for et, v, fmt in [
        ("Materialidad de la fase", _num(d.get("umbral")), PESOS),
        ("Piso de ruido", _num(d.get("trivialidad")), PESOS),
        ("Variación porcentual", _num(d.get("pct_variacion")), PORC),
        ("Cuentas comparadas", d["total_cuentas"], None),
        ("Cuentas seleccionadas", d["significativas"], None),
        ("Suma de lo no seleccionado", _num(d["residuo_no_seleccionado"]), PESOS),
    ]:
        h.cell(row=fila, column=1, value=et).font = _rotulo
        c = h.cell(row=fila, column=2, value=v)
        c.font = _normal
        if fmt:
            c.number_format = fmt
        fila += 1
    h.column_dimensions["A"].width = 32
    h.column_dimensions["B"].width = 22

    fila += 1
    if d.get("residuo_supera_umbral"):
        h.cell(row=fila, column=1,
               value="El conjunto no seleccionado supera la materialidad de la fase: "
                     "el alcance debe ampliarse o dejarse constancia de por qué se "
                     "acepta ese residuo.").font = Font(name="Calibri", size=10,
                                                        bold=True, color="B56A00")
        fila += 2

    _tabla(h, fila, ["Motivo de selección", "Cuentas"],
           [[k, v] for k, v in (d.get("por_motivo") or {}).items()],
           anchos=[32, 12])

    # -------------------------------------------------- 7. hallazgos
    h = _hoja(wb, "Hallazgos", "Hallazgos",
              "Detectados durante la carga y validación. Los no explicados no se "
              "fuerzan a cuadrar.")
    _tabla(h, 4, ["Severidad", "Tipo", "Cuenta", "Descripción", "Monto", "Fila origen"],
           [[x["severidad"], x["tipo"], x.get("codigo_puc") or "",
             x["descripcion"], _num(x.get("monto")), x.get("fila_origen")]
            for x in p["hallazgos"]] or [["—", "Sin hallazgos registrados", "", "", None, None]],
           anchos=[14, 22, 12, 70, 20, 12], formatos={5: PESOS}, principal=True)

    # ---------------------------------------------------- 8. riesgo
    h = _hoja(wb, "Riesgo", "Índice de riesgo",
              "Cada punto declara de dónde sale.")
    fila = _tabla(h, 4, ["Puntos", "Motivo", "Origen"],
        [[r["puntos"], r["motivo"], r["origen"]] for r in p["riesgo"]["detalle"]]
        or [[0, "Ningún control aportó puntos de riesgo", "—"]],
        anchos=[10, 70, 26], principal=True)
    h.cell(row=fila, column=1, value="TOTAL").font = _rotulo
    h.cell(row=fila, column=2,
           value=f"{p['riesgo']['puntos']} — {p['riesgo']['nivel']}").font = Font(
        name="Calibri", size=11, bold=True)

    # ---------------------------------------------------- 9. marcas
    h = _hoja(wb, "Marcas", "Marcas de auditoría (NIA 230)",
              "Una marca sin leyenda es un símbolo, no evidencia.")
    _tabla(h, 4, ["Marca", "Nombre", "NIA", "Procedimiento ejecutado",
                  "Contrastado contra", "¿Evidencia?"],
           [[m["marca"], m["nombre"], m["nia"], m["procedimiento"], m["contra"],
             "Sí" if m["evidencia"] else "No"] for m in p["marcas"]],
           anchos=[8, 28, 10, 62, 76, 13], principal=True)

    # ---------------------------------------------- 10. trazabilidad
    h = _hoja(wb, "Trazabilidad", "Trazabilidad",
              "Si una cifra del papel no se rastrea hasta un evento registrado, el "
              "papel no es auditable.")
    _tabla(h, 4, ["Fecha y hora", "Usuario", "Acción", "Éxito", "Origen", "Detalle"],
           [[e["creado_en"].strftime("%d/%m/%Y %H:%M:%S")
             if hasattr(e["creado_en"], "strftime") else str(e["creado_en"]),
             e.get("usuario") or "", e["accion"],
             "Sí" if e.get("exito") else "No",
             e.get("ip") or "", str(e.get("detalle") or "")]
            for e in p["trazabilidad"]["eventos"]],
           anchos=[20, 16, 26, 8, 18, 70], principal=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def nombre_archivo(p: dict) -> str:
    """Nombre que identifica el papel dentro de un archivo de auditoría
    sin tener que abrirlo."""
    id_ = p["identificacion"]
    corte = id_["fecha_corte"]
    corte = corte.isoformat() if isinstance(corte, date) else str(corte)
    nit = str(id_["nit"]).replace(" ", "")
    return f"PT_Variaciones_{nit}_{corte}_{id_['fase']}.xlsx"
