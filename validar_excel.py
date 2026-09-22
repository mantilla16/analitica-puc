"""
Pruebas de la lectura de archivos, sin base de datos.

    python validar_excel.py

`validar_motor.py` valida las reglas contables contra datos reales y por eso
necesita la base. Esto valida lo otro: cómo se interpreta un Excel antes de
que sus cifras signifiquen nada. Corre en cualquier parte, en dos segundos,
y los archivos se fabrican aquí mismo -- así las pruebas no dependen de
papeles de clientes que no pueden estar en el repositorio.

Cada caso es un defecto que ya ocurrió en producción. Están aquí para que no
vuelva a ocurrir en silencio, que es la única forma en que ocurrieron:

  · un ERP que reparte un mismo corte en varias hojas y solo se leía una;
  · dos columnas con el mismo rótulo, donde ganaba la última sin avisar;
  · una columna con datos y sin rótulo, que desaparecía del mapeo;
  · un archivo que declara mal su propio ancho y llegaba con una columna.
"""
from __future__ import annotations

import os
import tempfile
from collections import Counter

from openpyxl import Workbook

import excel as X

resultados: list[tuple[bool, str, str]] = []


def caso(nombre: str, esperado, obtenido, nota: str = "") -> None:
    ok = esperado == obtenido
    resultados.append((ok, nombre, nota if ok else
                       f"esperado {esperado!r}, obtenido {obtenido!r}. {nota}"))


def falla(nombre: str, fn, texto_esperado: str) -> None:
    """El caso pasa si `fn` se DETIENE mencionando `texto_esperado`.

    Que algo falle es la mitad del comportamiento correcto: un archivo al que
    le falta una hoja tiene que parar la carga, no seguir con lo que haya.
    """
    try:
        fn()
        resultados.append((False, nombre, "no se detuvo"))
    except ValueError as e:
        ok = texto_esperado.lower() in str(e).lower()
        resultados.append((ok, nombre, "" if ok else f"dijo: {str(e)[:120]}"))


# =====================================================================
# ARCHIVOS DE PRUEBA
# =====================================================================

CARPETA = tempfile.mkdtemp(prefix="validar_excel_")

MOV = ["Fecha", "Comprobante", "Secuencia", "Cuenta", "Nombre cuenta",
       "Nit", "Tercero", "Detalle", "Debitos", "Creditos"]
COLS_MOV = {"fecha": "Fecha", "num_doc": "Comprobante",
            "secuencia": "Secuencia", "codigo_puc": "Cuenta",
            "nombre_cuenta": "Nombre cuenta", "tercero_nit": "Nit",
            "tercero_nombre": "Tercero", "descripcion": "Detalle",
            "debito": "Debitos", "credito": "Creditos"}


def _guardar(wb, nombre: str) -> str:
    ruta = os.path.join(CARPETA, nombre)
    wb.save(ruta)
    return ruta


def archivo_en_tres_hojas() -> str:
    """Un corte repartido en tres pestañas, más una hoja de resumen."""
    wb = Workbook(); wb.remove(wb.active)
    for hoja, (desde, hasta) in (("MOV 1-500", (1, 500)),
                                 ("MOV 501-900", (501, 900)),
                                 ("MOV 901-1200", (901, 1200))):
        ws = wb.create_sheet(hoja)
        ws.append(["REPORTE DE MOVIMIENTOS"]); ws.append([])
        ws.append(MOV)
        for i in range(desde, hasta + 1):
            ws.append(["01/05/2026", f"CC-{i}", "1", "11050501", "CAJA GENERAL",
                       "900123456", "TERCERO SA", f"movimiento {i}", i * 100, 0])
    ws = wb.create_sheet("RESUMEN")
    ws.append(["Totales"]); ws.append([])
    ws.append(["Concepto", "Valor", "Otro", "Mas", "Aun mas"])
    ws.append(["Total debitos", 1000, 1, 2, 3])
    return _guardar(wb, "tres_hojas.xlsx")


def archivo_con_rotulos_repetidos() -> str:
    """Dos columnas DESCRIPCION, como el movimiento de KAIAK."""
    wb = Workbook(); ws = wb.active; ws.title = "MOV"
    ws.append(["Cuenta", "DESCRIPCION", "Comprobante", "DESCRIPCION",
               "Debitos", "Creditos"])
    ws.append(["11050501", "CAJA GENERAL", "CC-1", "PAGO A PROVEEDOR", 100, 0])
    return _guardar(wb, "rotulos_repetidos.xlsx")


def archivo_con_columna_sin_rotulo() -> str:
    """El nombre de la cuenta en una columna que nadie rotuló, como SEISA."""
    wb = Workbook(); ws = wb.active; ws.title = "Hoja 1"
    ws.append(["Cuentas", None, "2026/01", "Debitos", "Creditos", "2026/05"])
    # El rótulo va en la columna A, indentado con espacios para que
    # visualmente quede encima de la B. Así lo escribe SEISA, y por eso la
    # columna B se quedaba sin rótulo: para el Excel se ve bien, para quien
    # lo parsea la columna del nombre no existe.
    ws.append(["              Descripcion"])
    ws.append(["1", "ACTIVO", 100, 10, 5, 105])
    ws.append(["11", "DISPONIBLE", 50, 5, 2, 53])
    return _guardar(wb, "sin_rotulo.xlsx")


# =====================================================================
# UN CORTE REPARTIDO EN VARIAS HOJAS
# =====================================================================

def validar_varias_hojas() -> None:
    ruta = archivo_en_tres_hojas()
    TRES = ["MOV 1-500", "MOV 501-900", "MOV 901-1200"]

    def leer(perfil):
        return list(X.parsear(ruta, dict(perfil, formato_fecha="DD/MM/YYYY",
                                         columnas=COLS_MOV), "MOVIMIENTO"))

    caso("varias hojas · un perfil viejo sigue leyendo la suya",
         500, len(leer({"hoja": "MOV 1-500"})),
         "los perfiles guardados antes no pueden dejar de funcionar")

    filas = leer({"hojas": TRES})
    caso("varias hojas · se leen todas", 1200, len(filas))
    caso("varias hojas · cada fila sabe de dónde vino",
         {"MOV 1-500": 500, "MOV 501-900": 400, "MOV 901-1200": 300},
         dict(Counter(f["hoja"] for f in filas)))

    falla("varias hojas · una con otra estructura se rechaza",
          lambda: leer({"hojas": TRES + ["RESUMEN"]}),
          "RESUMEN")
    falla("varias hojas · una que no existe detiene la carga",
          lambda: leer({"hojas": ["MOV 1-500", "MOV DE OTRO MES"]}),
          "no tiene la(s) hoja(s)")


# =====================================================================
# DOS COLUMNAS CON EL MISMO ROTULO
# =====================================================================

def validar_rotulos_repetidos() -> None:
    ruta = archivo_con_rotulos_repetidos()
    h = X.inspeccionar(ruta)["hojas"][0]
    textos = [c["texto"] for c in h["columnas"]]

    caso("rótulos repetidos · se distinguen por su letra",
         ["Cuenta", "DESCRIPCION (col B)", "Comprobante",
          "DESCRIPCION (col D)", "Debitos", "Creditos"], textos,
         "si se llaman igual, el perfil no puede decir cuál quiere")

    filas = list(X.parsear(ruta, {
        "hojas": ["MOV"], "formato_fecha": "DD/MM/YYYY",
        "columnas": {"codigo_puc": "Cuenta",
                     "nombre_cuenta": "DESCRIPCION (col B)",
                     "num_doc": "Comprobante",
                     "descripcion": "DESCRIPCION (col D)",
                     "debito": "Debitos", "credito": "Creditos"},
    }, "MOVIMIENTO"))
    caso("rótulos repetidos · cada una trae lo suyo",
         ("CAJA GENERAL", "PAGO A PROVEEDOR"),
         (filas[0]["nombre_cuenta"], filas[0]["descripcion"]),
         "antes ganaba la última en silencio")


# =====================================================================
# UNA COLUMNA CON DATOS Y SIN ROTULO
# =====================================================================

def validar_columna_sin_rotulo() -> None:
    ruta = archivo_con_columna_sin_rotulo()
    h = X.inspeccionar(ruta)["hojas"][0]
    sin_rotulo = [c["texto"] for c in h["columnas"] if c.get("sin_encabezado")]

    caso("columna sin rótulo · se ofrece igual, nombrada por su letra",
         ["Columna B"], sin_rotulo,
         "sin esto desaparece del mapeo y el campo llega vacío")

    filas = list(X.parsear(ruta, {
        "hojas": ["Hoja 1"], "formato_fecha": "DD/MM/YYYY",
        "columnas": {"codigo_puc": "Cuentas", "nombre_cuenta": "Columna B",
                     "saldo_inicial": "2026/01", "debito": "Debitos",
                     "credito": "Creditos", "saldo_final": "2026/05"},
    }, "BALANCE"))
    caso("columna sin rótulo · se puede mapear y trae el dato",
         ["ACTIVO", "DISPONIBLE"], [f["nombre_cuenta"] for f in filas])


# =====================================================================
# UN ARCHIVO QUE MIENTE SOBRE SU PROPIO ANCHO
# =====================================================================

def validar_dimension_mal_declarada() -> None:
    """El caso de DOXA: el ERP declara `A1:A27591` teniendo 16 columnas.

    Se fabrica reescribiendo la dimensión dentro del .xlsx, que es un ZIP.
    Sin esto, openpyxl en modo read_only recorta cada fila a una columna.
    """
    import re, shutil, zipfile

    origen = archivo_con_rotulos_repetidos()
    roto = os.path.join(CARPETA, "dimension_mentirosa.xlsx")
    with zipfile.ZipFile(origen) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    for parte in list(partes):
        if parte.startswith("xl/worksheets/sheet"):
            xml = partes[parte].decode("utf8")
            # Ojo con el espacio: openpyxl escribe `<dimension ref="A1:F2" />`.
            # Sin contemplarlo, el reemplazo no ocurre y la prueba pasa sin
            # haber roto nada -- una prueba que no prueba nada es peor que
            # no tenerla, porque da confianza.
            xml, cambios = re.subn(r'<dimension ref="[^"]*"\s*/>',
                                   '<dimension ref="A1:A2"/>', xml)
            if cambios != 1:
                raise AssertionError(
                    f"No se pudo romper la dimension de {parte}: {cambios} cambios")
            partes[parte] = xml.encode("utf8")
    with zipfile.ZipFile(roto, "w", zipfile.ZIP_DEFLATED) as z:
        for parte, datos in partes.items():
            z.writestr(parte, datos)
    shutil.copystat(origen, roto)

    h = X.inspeccionar(roto)["hojas"][0]
    caso("dimensión mentirosa · se leen las columnas reales",
         6, len(h["columnas"]),
         "el archivo declara una sola columna; hay seis")
    caso("dimensión mentirosa · se declara la discrepancia",
         True, bool(h["dimension_mal_declarada"]),
         "es un dato sobre la fuente, no un detalle interno")


# =====================================================================

def main() -> int:
    validar_varias_hojas()
    validar_rotulos_repetidos()
    validar_columna_sin_rotulo()
    validar_dimension_mal_declarada()

    fallan = [r for r in resultados if not r[0]]
    for ok, nombre, nota in resultados:
        print(f"  {'PASA ' if ok else 'FALLA'}  {nombre}"
              + (f"\n            {nota}" if nota else ""))
    print(f"\n{len(resultados) - len(fallan)} de {len(resultados)} casos pasan.")
    if fallan:
        print(f"\nArchivos de prueba en {CARPETA}")
    return 1 if fallan else 0


if __name__ == "__main__":
    raise SystemExit(main())
